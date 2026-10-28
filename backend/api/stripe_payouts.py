"""
FastAPI Router for Stripe Payout Reconciliation & Approval
Endpoints for querying payouts, running read-only reconciliations, re-syncing with CRM,
approving and sending to Slack, and downloading enriched audit files.
"""

from fastapi import APIRouter, Query, HTTPException, Depends, Request
from fastapi.responses import FileResponse, StreamingResponse, Response
from pydantic import BaseModel
from typing import Optional, List, Dict, Any
import os
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.api.auth import get_current_user, require_company_access, require_super_admin
from core.stripe_accounts import ACCOUNTS, connected_companies, is_connected
from core.stripe_reconciliation import (
    STRIPE_COMPANY,
    ApprovalBlocked,
    classify_payout_line,
    list_stripe_payouts,
    reconcile_stripe_payout,
    resync_payout_with_crm,
    approve_and_send_to_slack,
    resend_payout_to_slack
)
from core.database import get_db_connection

async def _requested_company(request: Request) -> str:
    comp = request.query_params.get("company_id")
    if not comp and request.method in ("POST", "PUT", "PATCH"):
        try:
            body = await request.json()
            comp = body.get("company_id") if isinstance(body, dict) else None
        except Exception:
            comp = None
    return str(comp or STRIPE_COMPANY).strip().lower()


async def _stripe_account_access(request: Request):
    """Each charity's Stripe data is that charity's only: the request's charity (default Rethink) must have a
    connected Stripe account and the user must have access to it. 'all' is not a Stripe account."""
    user = getattr(request.state, "api_user", None) or await get_current_user(request, None, None)
    if request.url.path.endswith("/accounts"):
        return  # lists only the charities this user may see (labels, never keys)
    comp = await _requested_company(request)
    require_company_access(user, comp)
    if not is_connected(comp):
        raise HTTPException(status_code=404, detail=f"Stripe isn't connected for '{comp}' yet.")
    request.state.stripe_company = comp


router = APIRouter(prefix="/api/payouts/stripe", tags=["Stripe Payouts"], dependencies=[Depends(_stripe_account_access)])


def _account_company(requested: Optional[str]) -> str:
    comp = str(requested or STRIPE_COMPANY).strip().lower()
    if not is_connected(comp):
        raise HTTPException(status_code=404, detail=f"Stripe isn't connected for '{comp}' yet.")
    return comp


def _owned_payout(payout_id: str, company: str) -> None:
    """404 unless this payout was reconciled under this charity."""
    conn = get_db_connection(timeout=10.0)
    try:
        row = conn.execute("SELECT LOWER(COALESCE(company_id, 'rethink')) FROM stripe_payout_reports WHERE payout_id = ?",
                           (payout_id,)).fetchone()
    finally:
        conn.close()
    if not row or row[0] != company:
        raise HTTPException(status_code=404, detail="Payout not found for this charity.")


class ReconcileRequest(BaseModel):
    payout_id: str
    company_id: Optional[str] = "rethink"


class ApproveRequest(BaseModel):
    payout_id: str
    company_id: Optional[str] = "rethink"
    approved_by: Optional[str] = "Admin"


class ResendRequest(BaseModel):
    payout_id: str
    company_id: Optional[str] = "rethink"


class ResyncRequest(BaseModel):
    payout_id: str
    company_id: Optional[str] = "rethink"


@router.get("/list")
def get_stripe_payouts(limit: int = Query(25, ge=1, le=100), company_id: Optional[str] = Query("rethink")):
    """Lists Stripe payouts with their current CRM reconciliation and approval states."""
    try:
        payouts = list_stripe_payouts(limit=limit, company_id=_account_company(company_id))
        return {"payouts": payouts, "count": len(payouts)}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/reconcile")
def run_payout_reconciliation(req: ReconcileRequest, current_user = Depends(get_current_user)):
    """Runs read-only Stripe reconciliation for a payout, matching against CRM data."""
    try:
        summary = reconcile_stripe_payout(payout_id=req.payout_id, company_id=_account_company(req.company_id))
        return {"success": True, "summary": summary}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/resync")
def resync_payout_classifications(req: ResyncRequest, current_user = Depends(get_current_user)):
    """
    Re-syncs an already reconciled payout against current CRM classifications and donations.
    Prevents discrepancies if someone manually edited a donation code or classification in CRM.
    """
    try:
        comp = _account_company(req.company_id)
        _owned_payout(req.payout_id, comp)
        result = resync_payout_with_crm(payout_id=req.payout_id, company_id=comp)
        return {"success": True, "result": result}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/approve")
def approve_payout(req: ApproveRequest, current_user = Depends(get_current_user)):
    """Approves a payout and posts the verified financial breakdown + CSV report to Slack #rt-payouts."""
    try:
        approver = current_user.get("email") or "unknown"  # the signed-in account, not a name sent by the browser
        comp = _account_company(req.company_id)
        _owned_payout(req.payout_id, comp)
        result = approve_and_send_to_slack(payout_id=req.payout_id, approved_by=approver, company_id=comp)
        return result
    except ApprovalBlocked as e:
        raise HTTPException(status_code=409, detail=str(e))
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/resend")
def resend_to_slack(req: ResendRequest, current_user = Depends(get_current_user)):
    """Resends an already approved payout report to Slack #rt-payouts."""
    try:
        comp = _account_company(req.company_id)
        _owned_payout(req.payout_id, comp)
        result = resend_payout_to_slack(payout_id=req.payout_id, company_id=comp)
        return result
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/download")
def download_enriched_csv(payout_id: str = Query(...), company_id: Optional[str] = Query("rethink")):
    """Downloads the audit-ready enriched reconciliation CSV for a payout."""
    _owned_payout(payout_id, _account_company(company_id))
    conn = get_db_connection(timeout=10.0)
    try:
        cursor = conn.cursor()
        cursor.execute("SELECT csv_path FROM stripe_payout_reports WHERE payout_id = ?", (payout_id,))
        row = cursor.fetchone()
        if not row or not row[0] or not os.path.exists(row[0]):
            raise HTTPException(status_code=404, detail="Reconciliation CSV file not found.")
        file_path = row[0]
        return FileResponse(file_path, filename=os.path.basename(file_path), media_type="text/csv")
    finally:
        conn.close()


@router.get("/items")
def get_payout_items(
    payout_id: str = Query(...),
    source_type: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    company_id: Optional[str] = Query("rethink"),
):
    """Returns itemized transactions for a reconciled Stripe payout."""
    _owned_payout(payout_id, _account_company(company_id))
    conn = get_db_connection(timeout=10.0)
    try:
        cursor = conn.cursor()
        where_clauses = ["payout_id = ?"]
        params = [payout_id]

        if source_type and source_type != "ALL":
            where_clauses.append("source_type = ?")
            params.append(source_type.lower())

        where_sql = " AND ".join(where_clauses)
        cursor.execute(f"SELECT COUNT(*) FROM stripe_payout_items WHERE {where_sql}", params)
        total = cursor.fetchone()[0]

        offset = (page - 1) * page_size
        cursor.execute(f"""
            SELECT balance_transaction_id, charge_id, payment_intent_id, source_type, donation_id,
                   created_at, gross, fee, net, currency, donor_name, donor_email, project_name,
                   code, COALESCE(heading, department), COALESCE(subheading, office), department, office,
                   portfolio, country, programme_fund, fund_code, zakat_eligibility,
                   location_name, donor_phone, gift_aid, zakat_status, classified_by
            FROM stripe_payout_items
            WHERE {where_sql}
            ORDER BY created_at DESC
            LIMIT ? OFFSET ?
        """, params + [page_size, offset])

        cols = [
            "balance_transaction_id", "charge_id", "payment_intent_id", "source_type", "donation_id",
            "created_at", "gross", "fee", "net", "currency", "donor_name", "donor_email", "project_name",
            "code", "heading", "subheading", "department", "office",
            "portfolio", "country", "programme_fund", "fund_code", "zakat_eligibility",
            "location_name", "donor_phone", "gift_aid", "zakat_status", "classified_by"
        ]
        items = [dict(zip(cols, r)) for r in cursor.fetchall()]

        return {
            "payout_id": payout_id,
            "total": total,
            "page": page,
            "page_size": page_size,
            "total_pages": max(1, (total + page_size - 1) // page_size),
            "items": items
        }
    finally:
        conn.close()


@router.get("/export")
def export_stripe_payouts_excel(company_id: Optional[str] = "rethink"):
    company_id = _account_company(company_id)
    """
    Exports a comprehensive multi-sheet Excel (.xlsx) report containing all reconciled
    Stripe payouts, their financial summaries, and enriched transaction items.
    """
    conn = get_db_connection(timeout=10.0)
    try:
        cursor = conn.cursor()
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        header_fill = PatternFill(start_color="0E7490", end_color="0E7490", fill_type="solid")  # Dark Cyan
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        # -------------------------------------------------------------
        # Sheet 1: Stripe Payout Batches Summary
        # -------------------------------------------------------------
        ws1 = wb.create_sheet(title="Payout Batches")
        payouts_headers = [
            "Payout ID", "Arrival Date", "Status", "Currency", "Amount (£)",
            "Gross (£)", "Fees (£)", "Net (£)", "Transactions",
            "Reconciliation State", "Approval State", "Approved By", "Approved At"
        ]
        ws1.append(payouts_headers)
        for col_idx in range(1, len(payouts_headers) + 1):
            cell = ws1.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        cursor.execute("""
            SELECT payout_id, arrival_date, status, currency, amount,
                   total_gross, total_fees, total_net, total_transactions,
                   csv_path IS NOT NULL, approval_status, approved_by, approved_at
            FROM stripe_payout_reports
            WHERE LOWER(COALESCE(company_id, 'rethink')) = ?
            ORDER BY arrival_date DESC
        """, (company_id,))
        for r in cursor.fetchall():
            ws1.append([
                r[0], r[1], (r[2] or "").upper(), (r[3] or "GBP").upper(),
                float(r[4] or 0), float(r[5] or 0), float(r[6] or 0), float(r[7] or 0),
                int(r[8] or 0),
                "Reconciled" if r[9] else "Pending",
                (r[10] or "pending").title(),
                r[11] or "", r[12] or ""
            ])

        # -------------------------------------------------------------
        # Sheet 2: Reconciled Transactions
        # -------------------------------------------------------------
        ws2 = wb.create_sheet(title="Reconciled Transactions")
        tx_headers = [
            "Balance Transaction ID", "Charge ID", "Payout ID", "Date", "Source",
            "Donor Name", "Email", "Project / Description", "Gross (£)", "Fee (£)", "Net (£)",
            "Code", "Heading", "Sub-Heading", "Fund Code", "Zakat Eligibility", "Country"
        ]
        ws2.append(tx_headers)
        for col_idx in range(1, len(tx_headers) + 1):
            cell = ws2.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        cursor.execute("""
            SELECT balance_transaction_id, charge_id, payout_id, created_at, source_type,
                   donor_name, donor_email, project_name, gross, fee, net,
                   code, COALESCE(heading, department), COALESCE(subheading, office),
                   fund_code, zakat_eligibility, country
            FROM stripe_payout_items
            WHERE payout_id IN (SELECT payout_id FROM stripe_payout_reports WHERE LOWER(COALESCE(company_id, 'rethink')) = ?)
            ORDER BY created_at DESC
        """, (company_id,))
        for r in cursor.fetchall():
            ws2.append([
                r[0], r[1], r[2], r[3], (r[4] or "website").title(),
                r[5] or "", r[6] or "", r[7] or "",
                float(r[8] or 0), float(r[9] or 0), float(r[10] or 0),
                r[11] or "", r[12] or "", r[13] or "", r[14] or "", r[15] or "", r[16] or ""
            ])

        # Format column widths
        for ws in [ws1, ws2]:
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        output = io.BytesIO()
        wb.save(output)
        content_bytes = output.getvalue()

        filename = f"stripe_reconciliation_{company_id or 'rethink'}.xlsx"
        return Response(
            content=content_bytes,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    finally:
        conn.close()


class ClassifyLineRequest(BaseModel):
    balance_transaction_id: str
    code: str
    apply_to_project: bool = False
    company_id: Optional[str] = None


@router.post("/classify-line")
def classify_line(req: ClassifyLineRequest, current_user = Depends(get_current_user)):
    """Classifies one payout line (or, with apply_to_project, its project name in the matrix too)."""
    try:
        return classify_payout_line(req.balance_transaction_id, req.code, company_id=_account_company(req.company_id),
                                    apply_to_project=req.apply_to_project)
    except LookupError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.get("/needs-classification")
def needs_classification(limit: int = Query(500, ge=1, le=2000), company_id: Optional[str] = Query("rethink")):
    comp = _account_company(company_id)
    """Payout lines with no classification yet, or classified from a guess that a person should confirm."""
    conn = get_db_connection(timeout=10.0)
    try:
        rows = conn.execute("""
            SELECT i.balance_transaction_id, i.payout_id, r.arrival_date, i.created_at, i.source_type, i.project_name,
                   i.donor_name, i.donor_email, i.gross, i.net, i.currency, i.code, i.classified_by, i.charge_id
            FROM stripe_payout_items i JOIN stripe_payout_reports r ON r.payout_id = i.payout_id
            WHERE LOWER(COALESCE(r.company_id, 'rethink')) = ?
              AND (UPPER(i.code) = 'UNASSIGNED' OR i.classified_by = 'donor_history')
            ORDER BY i.created_at DESC LIMIT ?""", (comp, limit)).fetchall()
    finally:
        conn.close()
    cols = ["balance_transaction_id", "payout_id", "payout_arrival", "created_at", "source_type", "project_name",
            "donor_name", "donor_email", "gross", "net", "currency", "code", "classified_by", "charge_id"]
    items = [dict(zip(cols, r)) for r in rows]
    for it in items:
        it["reason"] = "Guessed from the donor's previous gift" if it["classified_by"] == "donor_history" else "No matching classification"
    return {"items": items, "count": len(items)}


@router.get("/ledger-summary")
def ledger_summary(company_id: Optional[str] = Query("rethink")):
    """Live position of the Stripe account from the ledger (kept current by the sync service every minute)."""
    comp = _account_company(company_id)
    from core.stripe_ledger import _skey
    conn = get_db_connection(timeout=10.0)
    try:
        has = conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='stripe_balance_transactions'").fetchone()
        if not has:
            return {"available": False}
        q = lambda sql, *a: conn.execute(sql, a).fetchone()  # noqa: E731
        unpaid = q("""SELECT COUNT(*), COALESCE(ROUND(SUM(net), 2), 0) FROM stripe_balance_transactions
                      WHERE company_id = ? AND payout_id IS NULL AND type != 'payout'""", comp)
        gifts_30 = q("""SELECT COUNT(*), COALESCE(ROUND(SUM(amount), 2), 0) FROM stripe_balance_transactions
                        WHERE company_id = ? AND type IN ('charge', 'payment') AND created_at >= datetime('now', '-30 days')""", comp)
        refunds_30 = q("""SELECT COUNT(*), COALESCE(ROUND(SUM(amount), 2), 0) FROM stripe_balance_transactions
                          WHERE company_id = ? AND (type LIKE '%refund%' OR reporting_category LIKE 'dispute%')
                            AND created_at >= datetime('now', '-30 days')""", comp)
        disputed = q("""SELECT COUNT(*) FROM stripe_balance_transactions WHERE company_id = ? AND disputed = 1""", comp)
        last = q("SELECT updated_at FROM stripe_sync_state WHERE key = ?", _skey(comp, "recent_cursor"))
        backfill = q("SELECT value FROM stripe_sync_state WHERE key = ?", _skey(comp, "backfill_done"))
        total = q("SELECT COUNT(*), MIN(created_at) FROM stripe_balance_transactions WHERE company_id = ?", comp)
        open_items = q("""SELECT COUNT(*) FROM stripe_payout_items i JOIN stripe_payout_reports r ON r.payout_id = i.payout_id
                          WHERE LOWER(COALESCE(r.company_id,'rethink')) = ? AND (UPPER(i.code) = 'UNASSIGNED' OR i.classified_by = 'donor_history')""",
                       comp)
    finally:
        conn.close()
    return {
        "available": True,
        "awaiting_payout": {"count": unpaid[0], "net": unpaid[1]},
        "gifts_30_days": {"count": gifts_30[0], "amount": gifts_30[1]},
        "refunds_30_days": {"count": refunds_30[0], "amount": refunds_30[1]},
        "disputed_charges": disputed[0],
        "needs_classification": open_items[0],
        "last_synced_at": last[0] if last else None,
        "history_complete": bool(backfill and backfill[0] == "1"),
        "movements_recorded": total[0],
        "recorded_since": total[1],
    }


@router.get("/accounts")
def stripe_accounts(request: Request):
    """Charities with a connected Stripe account that this user may see (labels only - never keys)."""
    user = getattr(request.state, "api_user", None) or {}
    from backend.api.auth import user_company_ids
    allowed = user_company_ids(user)
    return {"accounts": [{"company_id": c, "label": ACCOUNTS[c]["label"]} for c in connected_companies()
                         if allowed is None or c in allowed]}
